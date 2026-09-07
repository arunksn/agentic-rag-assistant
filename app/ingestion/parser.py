from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".csv",
    ".txt",
    ".json",
    ".md",
}


def parse_pdf(path: Path) -> list[Document]:
    """Parse a PDF into one Document per page."""

    reader = PdfReader(str(path))
    documents: list[Document] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "page": page_number,
                },
            )
        )

    return documents


def parse_docx(path: Path) -> list[Document]:
    """Extract paragraph text from a DOCX document."""

    document = DocxDocument(str(path))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return [
        Document(
            page_content="\n".join(paragraphs),
            metadata={},
        )
    ]


def parse_xlsx(path: Path) -> list[Document]:
    """Parse an XLSX workbook into one Document per worksheet."""

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    documents: list[Document] = []

    for worksheet in workbook.worksheets:
        rows: list[str] = []

        for row in worksheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

        documents.append(
            Document(
                page_content="\n".join(rows),
                metadata={
                    "sheet": worksheet.title,
                },
            )
        )

    return documents


def parse_csv(path: Path) -> list[Document]:
    """Parse a CSV file into a text representation."""

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.reader(file)
        rows = [
            " | ".join(cell.strip() for cell in row)
            for row in reader
        ]

    return [
        Document(
            page_content="\n".join(row for row in rows if row.strip()),
            metadata={},
        )
    ]


def parse_json(path: Path) -> list[Document]:
    """Parse JSON into a readable text representation."""

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    return [
        Document(
            page_content=json.dumps(
                data,
                indent=2,
                ensure_ascii=False,
            ),
            metadata={},
        )
    ]


def parse_text(path: Path) -> list[Document]:
    """Parse plain-text or Markdown documents."""

    return [
        Document(
            page_content=path.read_text(
                encoding="utf-8",
            ),
            metadata={},
        )
    ]


def parse_file(path: Path) -> list[Document]:
    """Dispatch a file to its appropriate parser."""

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        return parse_pdf(path)

    if extension == ".docx":
        return parse_docx(path)

    if extension == ".xlsx":
        return parse_xlsx(path)

    if extension == ".csv":
        return parse_csv(path)

    if extension == ".json":
        return parse_json(path)

    return parse_text(path)